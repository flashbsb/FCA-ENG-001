#!/usr/bin/env python3
"""
Script de extração e análise de lote — FCA-ENG-001 v2.0
Extrai texto de todos os arquivos de um lote e retorna como dicionário.
Usado pela IA para alimentar o processo de análise RTI.
"""
import sys, os, json, zipfile, re, xml.etree.ElementTree as ET

def get_workspace_root():
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def extract_docx_text(fpath):
    try:
        z = zipfile.ZipFile(fpath)
        doc_xml = z.read('word/document.xml')
        tree = ET.fromstring(doc_xml)
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        paras = []
        for p in tree.findall('.//w:p', ns):
            texts = [r.text or '' for r in p.findall('.//w:t', ns)]
            t = ''.join(texts).strip()
            if t:
                paras.append(t)
        return '\n'.join(paras)
    except Exception as e:
        return f"[ERRO ao ler .docx: {e}]"

def extract_pdf_text(fpath):
    try:
        import pypdf
        reader = pypdf.PdfReader(fpath)
        pages = []
        for i, page in enumerate(reader.pages):
            t = page.extract_text()
            if t:
                pages.append(f"[Página {i+1}]\n{t}")
        return '\n\n'.join(pages)
    except ImportError:
        return "[pypdf não disponível]"
    except Exception as e:
        return f"[ERRO ao ler .pdf: {e}]"

def extract_xlsx_text(fpath):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        result = []
        for name in wb.sheetnames:
            ws = wb[name]
            result.append(f"\n## Aba: {name}\n")
            for row in ws.iter_rows(max_row=500, values_only=True):
                row_vals = [str(c) if c is not None else '' for c in row]
                if any(v.strip() for v in row_vals):
                    result.append(' | '.join(row_vals))
        wb.close()
        return '\n'.join(result)
    except ImportError:
        return "[openpyxl não disponível]"
    except Exception as e:
        return f"[ERRO ao ler .xlsx: {e}]"

def extract_text_file(fpath):
    try:
        with open(fpath, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read(500000)
    except Exception as e:
        return f"[ERRO: {e}]"

def extract_lote(lote_path):
    """Extrai todo o conteúdo de um lote para análise pela IA."""
    result = {'lote': os.path.basename(lote_path), 'files': {}}
    
    ext_map = {
        '.docx': extract_docx_text,
        '.pdf': extract_pdf_text,
        '.xlsx': extract_xlsx_text,
        '.xls': extract_xlsx_text,
        '.csv': extract_text_file,
        '.txt': extract_text_file,
        '.md': extract_text_file,
        '.json': extract_text_file,
        '.yaml': extract_text_file,
        '.yml': extract_text_file,
    }
    
    for dirpath, _, filenames in os.walk(lote_path):
        for fname in sorted(filenames):
            if fname.startswith('.') or fname in ('README.md', 'decisoes.yaml', '.gitkeep'):
                continue
            fpath = os.path.join(dirpath, fname)
            ext = os.path.splitext(fname)[1].lower()
            rel = os.path.relpath(fpath, lote_path)
            
            if ext in ext_map:
                print(f"  Extraindo: {rel}...", file=sys.stderr)
                content = ext_map[ext](fpath)
                result['files'][rel] = {
                    'name': fname,
                    'ext': ext,
                    'size': os.path.getsize(fpath),
                    'content': content,
                    'content_length': len(content)
                }
            else:
                result['files'][rel] = {
                    'name': fname,
                    'ext': ext,
                    'size': os.path.getsize(fpath),
                    'content': None,
                    'content_length': 0,
                    'note': f"Formato {ext} não extraível (imagem/binário)"
                }
    
    return result

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Uso: python3 extract_lote.py <lote_dir>", file=sys.stderr)
        sys.exit(1)
    
    lote_path = sys.argv[1]
    if not os.path.isabs(lote_path):
        lote_path = os.path.join(get_workspace_root(), lote_path)
    
    print(f"Extraindo lote: {lote_path}", file=sys.stderr)
    data = extract_lote(lote_path)
    
    # Gera sumário para análise
    print(f"\n{'='*60}", file=sys.stderr)
    print(f"Lote: {data['lote']}", file=sys.stderr)
    for rel, fdata in data['files'].items():
        status = f"{fdata['content_length']} chars" if fdata.get('content') else fdata.get('note', 'N/A')
        print(f"  {rel}: {status}", file=sys.stderr)
    
    json.dump(data, sys.stdout, ensure_ascii=False, indent=2)
